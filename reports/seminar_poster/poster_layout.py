"""Build the current A1 poster in the user's stacked-section reference layout."""
import csv
import json
from pathlib import Path
from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
W,H=594/25.4,841/25.4
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
s=prs.slides.add_slide(prs.slide_layouts[6])
BLUE='075D93';CYAN='39BED1';GREEN='226D4A';PALE='EEF7F3';INK='243747';MUTED='526779';WHITE='FFFFFF';LINE='D8E8ED'
s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)


def shape(x,y,w,h,fill=None,line=None,kind=MSO_SHAPE.RECTANGLE,width=1.4):
    item=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:item.fill.solid();item.fill.fore_color.rgb=RGBColor.from_string(fill)
    else:item.fill.background()
    if line:item.line.color.rgb=RGBColor.from_string(line);item.line.width=Pt(width)
    else:item.line.fill.background()
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:item.adjustments[0]=.08
    return item


def text(value,x,y,w,h,size=27,color=INK,bold=False,align=PP_ALIGN.LEFT):
    item=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=item.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.04);tf.margin_top=tf.margin_bottom=Inches(.015)
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos'
        p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.alignment=align;p.space_after=Pt(5);p.space_before=Pt(0)
    return item


def section(title,y,h):
    shape(.7,y,W-1.4,h,WHITE,CYAN,MSO_SHAPE.ROUNDED_RECTANGLE,1.8)
    text(title,1.2,y+.16,W-2.4,.55,30,BLUE,True,PP_ALIGN.CENTER)


def leaf(x,y,scale=1,color=GREEN):
    sh=shape(x,y,.72*scale,1.15*scale,color,None,MSO_SHAPE.OVAL);sh.rotation=37
    vein=shape(x+.39*scale,y+.25*scale,.035*scale,.78*scale,'C8E6D3');vein.rotation=37


def plant(x,y,scale=1):
    shape(x+.75*scale,y+.7*scale,.07*scale,1.2*scale,GREEN)
    leaf(x+.2*scale,y+.48*scale,.75*scale)
    leaf(x+.85*scale,y+.18*scale,.75*scale)
    shape(x+.2*scale,y+1.8*scale,1.2*scale,.5*scale,BLUE,None,MSO_SHAPE.TRAPEZOID)


def arrow(x,y,w=.36,h=.35):
    shape(x,y,w,h,CYAN,None,MSO_SHAPE.DOWN_ARROW)
