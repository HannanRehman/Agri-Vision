"""Build a native, editable A1 seminar poster from recorded project results."""
from pathlib import Path
import csv
import json
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width = Inches(594 / 25.4)
prs.slide_height = Inches(841 / 25.4)
s = prs.slides.add_slide(prs.slide_layouts[6])
GREEN = '173F32'
SAGE = 'E7EDE3'
CREAM = 'F7F5ED'
INK = '233B32'
MUTED = '53685F'
AMBER = 'B97722'
WHITE = 'FFFFFF'
s.background.fill.solid()
s.background.fill.fore_color.rgb = RGBColor.from_string(CREAM)

def box(x, y, w, h, fill, line=None, radius=False):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                            Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line)
    else:
        sh.line.fill.background()
    if radius:
        sh.adjustments[0] = .08
    return sh

def txt(text, x, y, w, h, size=27, color=INK, bold=False, align=PP_ALIGN.LEFT):
    sh = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.02)
    tf.margin_top = tf.margin_bottom = Inches(.02)
    for i, line in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = 'Aptos'
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.alignment = align
        p.space_after = Pt(9)
    return sh

def heading(num, title, x, y, w=10.2):
    txt(num, x, y, .6, .6, 28, AMBER, True)
    txt(title, x+.7, y, w-.7, .65, 32, GREEN, True)
    box(x, y+.78, w, .035, GREEN)

# Header illustration is decorative; it is never presented as an experiment.
s.shapes.add_picture(str(OUT/'assets/crop_illustration.png'), Inches(10.4), Inches(.3), width=Inches(12.7))
txt('FINAL YEAR PROJECT  /  SEMINAR', .95, .65, 10.2, .6, 24, AMBER, True)
txt('AGRI\nVISION', .9, 1.15, 10.5, 4.3, 118, GREEN, True)
txt('Federated learning for\ncrop-disease detection', .95, 5.65, 10.6, 1.25, 34, GREEN, True)
txt('Hannan Rehman  •  Fuzail Talpur  •  Ali', .95, 7.15, 21.4, .55, 27, GREEN, True)
txt('Supervisor: Dr Rabeea Jaffari', .95, 7.8, 21.4, .55, 27, MUTED)
box(.0, 8.65, 23.39, .95, GREEN)
txt('BANANA   •   CAULIFLOWER   •   CHILI   •   GROUNDNUT   •   RADISH', .8, 8.85, 21.8, .5, 27, WHITE, True, PP_ALIGN.CENTER)

txt('FIVE CROPS. THIRTY DETECTION CLASSES.', .95, 10.05, 21.5, .55, 31, GREEN, True)
patterns = ['*sigatoka*', '*Black-Rot*', '*Anthracnose*', '*Early-Leaf-Spot*', '*Black-Leaf-Spot*']
names = ['Banana', 'Cauliflower', 'Chili', 'Groundnut', 'Radish']
photos = []
for i, (pattern, name) in enumerate(zip(patterns, names)):
    p = sorted((ROOT/'original_dataset/test/images').glob(pattern))[0]
    photos.append(str(p.relative_to(ROOT)))
    x = .95 + i*4.36
    # Native PPTX image cropping; original photographs remain unchanged.
    sh = s.shapes.add_picture(str(p), Inches(x), Inches(10.85), width=Inches(4.04), height=Inches(2.05))
    with Image.open(p) as photo:
        ratio = photo.width / photo.height
    target_ratio = 4.04 / 2.05
    if ratio > target_ratio:
        sh.crop_left = sh.crop_right = (1-target_ratio/ratio)/2
    else:
        sh.crop_top = sh.crop_bottom = (1-ratio/target_ratio)/2
    txt(name, x, 13.04, 4.04, .5, 26, GREEN, True, PP_ALIGN.CENTER)
txt('Examples from the labelled dataset; these photographs are not model predictions.', .95, 13.7, 21.4, .5, 22, MUTED)

L, R = .95, 12.05
heading('01', 'THE CHALLENGE', L, 14.55)
txt('Visible disease symptoms can be small, similar across classes, and unevenly represented in crop datasets. Agri Vision explores automated symptom localisation while training a shared model across separate client partitions.', L, 15.65, 10.2, 2.3, 28)
heading('02', 'OUR APPROACH', L, 18.15)
txt('YOLOv8n detects symptom regions and predicts their classes. Three clients train locally from a shared checkpoint; sample-weighted FedAvg combines their model weights.', L, 19.25, 10.2, 1.9, 28)

# Fully editable federated-learning diagram.
for i in range(3):
    x = L + i*3.46
    box(x, 21.45, 3.15, 1.15, SAGE, radius=True)
    txt(f'CLIENT {i+1}\nLocal training', x+.12, 21.59, 2.91, .92, 23, GREEN, True, PP_ALIGN.CENTER)
    arrow = s.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(x+1.35), Inches(22.73), Inches(.45), Inches(.6))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = RGBColor.from_string(AMBER); arrow.line.fill.background()
box(L, 23.55, 10.2, .92, GREEN, radius=True)
txt('SAMPLE-WEIGHTED FEDAVG  →  GLOBAL MODEL', L+.2, 23.77, 9.8, .5, 25, WHITE, True, PP_ALIGN.CENTER)
txt('Repeat local training → aggregate → validate', L, 24.7, 10.2, .55, 25, MUTED, False, PP_ALIGN.CENTER)
txt('Prototype: sequential three-client simulation on one Tesla T4. This is not a deployed multi-device system.', L, 25.4, 10.2, 1.0, 23, MUTED)

heading('03', 'HELD-OUT TEST RESULTS', R, 14.55)
txt('Round 35  •  All 30 classes  •  2,191 test images', R, 15.6, 10.2, .6, 25, MUTED)
scores = json.loads((ROOT/'reports/round_035_non_chili_summary.json').read_text())['groups']['all_30_classes']
for i, (label, key) in enumerate([('Precision','precision'), ('Recall','recall'), ('mAP@50','mAP50'), ('mAP@50–95','mAP50_95')]):
    x, y = R+(i%2)*5.2, 16.6+(i//2)*1.8
    box(x, y, 4.97, 1.6, SAGE, radius=True)
    txt(f'{scores[key]*100:.2f}%', x+.2, y+.12, 4.57, .84, 47, GREEN, True)
    txt(label, x+.2, y+1.02, 4.57, .47, 25, MUTED)
txt('Macro averages across classes. Precision and recall use the evaluation operating point selected over all 30 classes.', R, 20.32, 10.2, .92, 22, MUTED)
heading('04', 'TRAINING TREND', R, 21.55)
rows = []
for f in ['rounds_026_035_global_metrics.csv', 'rounds_036_040_global_metrics.csv']:
    with (ROOT/'reports'/f).open(newline='') as fh:
        rows.extend(csv.DictReader(fh))
data = CategoryChartData()
data.categories = [r['round'] for r in rows]
data.add_series('mAP@50', [float(r['mAP50'])*100 for r in rows])
chart = s.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(R-.2), Inches(22.5), Inches(10.55), Inches(2.8), data).chart
chart.has_legend = False
chart.chart_style = 13
chart.value_axis.minimum_scale = 58
chart.value_axis.maximum_scale = 63
chart.value_axis.major_unit = 1
chart.value_axis.tick_labels.number_format = '0"%"'
for ax in [chart.value_axis, chart.category_axis]:
    ax.tick_labels.font.name = 'Aptos'; ax.tick_labels.font.size = Pt(20)
chart.category_axis.tick_label_spacing = 2
series = chart.series[0]
series.format.line.color.rgb = RGBColor.from_string(GREEN)
series.format.line.width = Pt(3.5)
txt('Rounds 26–40 · Client-weighted validation mAP@50', R, 25.5, 10.2, .5, 22, MUTED)

box(.95, 27.0, 21.3, 4.12, GREEN, radius=True)
txt('WHAT THE RESULTS SHOW', 1.35, 27.35, 20.45, .65, 34, WHITE, True)
txt('The latest checkpoint was not the strongest.', 1.35, 28.2, 20.45, .6, 30, 'DCE8D2', True)
txt('Among rounds 26–40, validation mAP@50 peaked at 61.22% in round 33, then reached 59.55% in round 40. Validation mAP@50–95 peaked at 40.61% in round 30. These validation scores are separate from the held-out test results above.', 1.35, 29.0, 20.45, 1.55, 27, WHITE)
txt('NEXT: review weak-class labels and data balance; select checkpoints on validation performance; run a controlled Faster R-CNN comparison.', .95, 31.45, 21.3, .95, 25, GREEN, True)
txt('DATA SPLIT  15,512 train  /  4,374 validation  /  2,191 test     •     MODEL  YOLOv8n + FedAvg', .95, 32.58, 21.3, .4, 20, MUTED)

prs.core_properties.title = 'Agri Vision — FYP Seminar Poster'
prs.core_properties.subject = 'A1 portrait; crop-disease detection with three-client federated learning'
prs.core_properties.author = 'Hannan Rehman, Fuzail Talpur, Ali'
prs.save(OUT/'Agri_Vision_Seminar_Poster_A1.pptx')
(OUT/'sources.json').write_text(json.dumps({'test_metrics':'reports/round_035_non_chili_summary.json', 'validation_metrics':['reports/rounds_026_035_global_metrics.csv','reports/rounds_036_040_global_metrics.csv'], 'dataset_examples':photos, 'header':'AI-generated botanical illustration; decorative only', 'format':'A1 portrait, 594 × 841 mm'}, indent=2))
print(OUT/'Agri_Vision_Seminar_Poster_A1.pptx')
