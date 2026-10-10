"""Build the current A1 poster in the user's stacked-section reference layout."""
import csv
import json
from pathlib import Path
from PIL import Image

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).resolve().parent
W,H=594/25.4,841/25.4
prs=Presentation();prs.slide_width=Inches(W);prs.slide_height=Inches(H)
s=prs.slides.add_slide(prs.slide_layouts[6])
BLUE='075D93';CYAN='39BED1';GREEN='226D4A';PALE='EEF7F3';INK='243747';MUTED='526779';WHITE='FFFFFF';LINE='D8E8ED'
s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)


def shape(x,y,w,h,fill=None,line=None,kind=MSO_SHAPE.RECTANGLE,width=1.4):
    item=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    if fill:item.fill.solid();item.fill.fore_color.rgb=RGBColor.from_string(fill)
    else:item.fill.background()
    if line:item.line.color.rgb=RGBColor.from_string(line);item.line.width=Pt(width)
    else:item.line.fill.background()
    if kind==MSO_SHAPE.ROUNDED_RECTANGLE:item.adjustments[0]=.08
    return item


def text(value,x,y,w,h,size=27,color=INK,bold=False,align=PP_ALIGN.LEFT):
    item=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=item.text_frame
    tf.word_wrap=True;tf.margin_left=tf.margin_right=Inches(.04);tf.margin_top=tf.margin_bottom=Inches(.015)
    for i,line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Aptos'
        p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
        p.alignment=align;p.space_after=Pt(5);p.space_before=Pt(0)
    return item


def section(title,y,h):
    shape(.7,y,W-1.4,h,WHITE,CYAN,MSO_SHAPE.ROUNDED_RECTANGLE,1.8)
    text(title,1.2,y+.16,W-2.4,.55,30,BLUE,True,PP_ALIGN.CENTER)


def leaf(x,y,scale=1,color=GREEN):
    sh=shape(x,y,.72*scale,1.15*scale,color,None,MSO_SHAPE.OVAL);sh.rotation=37
    vein=shape(x+.39*scale,y+.25*scale,.035*scale,.78*scale,'C8E6D3');vein.rotation=37


def plant(x,y,scale=1):
    shape(x+.75*scale,y+.7*scale,.07*scale,1.2*scale,GREEN)
    leaf(x+.2*scale,y+.48*scale,.75*scale)
    leaf(x+.85*scale,y+.18*scale,.75*scale)
    shape(x+.2*scale,y+1.8*scale,1.2*scale,.5*scale,BLUE,None,MSO_SHAPE.TRAPEZOID)


def arrow(x,y,w=.36,h=.35):
    shape(x,y,w,h,CYAN,None,MSO_SHAPE.DOWN_ARROW)


# White institutional-header area: use crop identities rather than invented logos.
text('FINAL YEAR PROJECT  /  SEMINAR',1.05,.48,9.1,.48,25,BLUE,True)
text('Federated learning for crop-disease detection',1.05,1.1,11.0,.65,26,MUTED)
for i,name in enumerate(['BANANA','CAULIFLOWER','GROUNDNUT','RADISH']):
    x=12.25+i*2.47
    leaf(x+.82,.39,.6)
    text(name,x,1.27,2.42,.35,16,GREEN,True,PP_ALIGN.CENTER)
shape(0,2.0,W,2.03,BLUE)
text('AGRI VISION',1.0,2.08,W-2,1.27,73,WHITE,True,PP_ALIGN.CENTER)
text('Detecting crop symptoms. Comparing lightweight one-stage detectors.',1.0,3.49,W-2,.45,25,WHITE,False,PP_ALIGN.CENTER)
text('Hannan Rehman  •  Fuzail Talpur  •  Ali',1.0,4.19,12.6,.52,26,BLUE,True)
text('Supervisor: Dr Rabeea Jaffari',13.0,4.19,9.35,.52,26,MUTED,False,PP_ALIGN.RIGHT)

section('ABSTRACT',5.0,2.85)
plant(1.45,5.67,.87)
text('Agri Vision investigates crop-disease and pest-symptom detection using bounding boxes and disease labels. '
     'A three-client federated-learning simulation trains a shared detector across separate data partitions. '
     'The revised study covers banana, cauliflower, groundnut and radish across 24 classes. '
     'SSDLite, SSD–MobileNetV3 and EfficientDet-D0 are being compared with the retained historical YOLOv8n baseline.',
     4.0,5.82,17.75,1.82,26)

section('PROBLEM STATEMENT',8.10,2.55)
text('• Small lesions and similar symptoms make visual diagnosis difficult.\n'
     '• Uneven class representation and variable image quality affect detection.\n'
     '• A useful detector must balance precision, recall, localisation and computing cost.',
     1.35,8.89,17.4,1.51,26)
shape(19.45,8.93,1.3,1.3,PALE,None,MSO_SHAPE.OVAL)
leaf(19.81,9.02,.85)
for dx,dy in [(.22,.36),(.49,.61),(.33,.76)]:shape(19.81+dx,9.02+dy,.14,.14,'AE6B38',None,MSO_SHAPE.OVAL)
shape(20.51,10.02,.5,.08,BLUE).rotation=45

section('AIM AND OBJECTIVES',10.90,2.70)
plant(1.4,11.53,.76)
text('Develop and assess an efficient federated crop-symptom detector.\n'
     '• Use one verified 24-class dataset with fixed train, validation and test splits.\n'
     '• Compare four complete one-stage detectors through separate experiments.\n'
     '• Measure precision, recall, mAP, inference speed and model size.',
     4.0,11.63,17.75,1.92,25)

section('METHODOLOGY',13.85,5.30)
text('Shared 24-class dataset  →  fixed three-client partitions  →  independent detector experiments',
     1.1,14.58,W-2.2,.46,25,MUTED,False,PP_ALIGN.CENTER)
for i in range(3):
    x=2.0+i*7.0
    shape(x,15.25,5.3,.88,PALE,LINE,MSO_SHAPE.ROUNDED_RECTANGLE)
    text(f'CLIENT {i+1}  •  LOCAL TRAINING',x+.15,15.45,5.0,.45,24,GREEN,True,PP_ALIGN.CENTER)
    arrow(x+2.46,16.17)
shape(2.0,16.62,19.3,.67,BLUE,None,MSO_SHAPE.ROUNDED_RECTANGLE)
text('TRAINING-IMAGE-WEIGHTED FEDAVG  →  GLOBAL DETECTOR',2.25,16.73,18.8,.45,25,WHITE,True,PP_ALIGN.CENTER)
for i,name in enumerate(['YOLOv8n','SSDLite','SSD–MobileNetV3','EfficientDet-D0']):
    x=1.4+i*5.35
    shape(x,17.72,5.05,.64,PALE,LINE,MSO_SHAPE.ROUNDED_RECTANGLE)
    text(name,x+.08,17.82,4.89,.43,23,GREEN,True,PP_ALIGN.CENTER)
text('Train locally  →  aggregate  →  validate  →  save best/latest checkpoint  →  repeat',
     1.25,18.42,W-2.5,.45,24,BLUE,True,PP_ALIGN.CENTER)
text('Each architecture has its own weights and run. Clients are simulated sequentially on one GPU; this is not a deployed multi-device network.',
     1.3,18.91,W-2.6,.23,14,MUTED,False,PP_ALIGN.CENTER)

section('RESULTS AND CURRENT PROGRESS',19.45,8.85)
text('PROVISIONAL 24-CLASS VALIDATION',1.35,20.26,11.1,.53,25,GREEN,True)
rows=list(csv.DictReader((ROOT/'reports/comparison/main_run_review_20261010/validation_summary.csv').open()))
by_model={row['model']:row for row in rows}
headers=['Detector','Round','Precision','Recall','mAP50–95']
widths=[3.55,1.12,2.04,1.83,2.22]
x0,y0=1.35,21.05
xx=x0
for label,width in zip(headers,widths):
    shape(xx,y0,width,.58,BLUE)
    text(label,xx+.08,y0+.12,width-.16,.37,21,WHITE,True,PP_ALIGN.LEFT if label=='Detector' else PP_ALIGN.CENTER)
    xx+=width
for i,(model,label) in enumerate([('ssdlite','SSDLite'),('ssd_mobilenetv3','SSD–MobileNetV3'),('efficientdet_d0','EfficientDet-D0')]):
    row=by_model[model];values=[label,row['best_round'],f"{float(row['precision'])*100:.2f}%",f"{float(row['recall'])*100:.2f}%",f"{float(row['map_50_95'])*100:.2f}%"]
    yy=y0+.58+i*.67;xx=x0
    for j,(value,width) in enumerate(zip(values,widths)):
        shape(xx,yy,width,.67,PALE if i%2==0 else 'F6FAFC',WHITE)
        text(value,xx+.08,yy+.18,width-.16,.36,21,INK,j==0,PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER)
        xx+=width
text('Same 4,157-image pooled validation split. P/R: confidence 0.25, IoU 0.50. '
     'EfficientDet has completed only two rounds; these scores do not establish a final winner.',
     1.35,23.83,10.76,1.14,23,MUTED)

historical=list(csv.DictReader((ROOT/'reports/rounds_036_040_global_metrics.csv').open()))[-1]
shape(1.35,25.12,10.76,1.58,'EAF2FA',None,MSO_SHAPE.ROUNDED_RECTANGLE)
text('HISTORICAL YOLOv8n  •  ROUND 40',1.58,25.31,10.3,.41,24,BLUE,True)
text(f"Precision {float(historical['precision'])*100:.2f}%    Recall {float(historical['recall'])*100:.2f}%\n"
     f"mAP50 {float(historical['mAP50'])*100:.2f}%    mAP50–95 {float(historical['mAP50_95'])*100:.2f}%",
     1.58,25.76,10.3,.91,24,BLUE)
text('30-class training; validation-image-weighted native client summaries. '
     'A common 24-class evaluation of this checkpoint is pending; the scores above are not directly comparable.',
     1.35,26.89,10.76,1.12,21,MUTED)

# Four editable phone-style cards, matching the reference's 2×2 result-image grid.
# Photographs/boxes are ground-truth examples, never claimed as model predictions.
names=json.loads((ROOT/'manifests/non_chili_v1/class_map.json').read_text())['names']
with (ROOT/'manifests/non_chili_v1/retained.csv').open(newline='',encoding='utf-8') as stream:manifest=list(csv.DictReader(stream))
photos=[]
for i,(crop,disease) in enumerate([('Banana','sigatoka'),('Cauliflower','black_rot'),('Groundnut','early_leaf_spot'),('Radish','black_leaf_spot')]):
    class_id=next(int(k) for k,v in names.items() if v.lower()==f'{crop}_{disease}'.lower())
    candidates=[]
    for row in manifest:
        if row['split']!='valid':continue
        annotations=[line.split() for line in (ROOT/row['label']).read_text().splitlines() if line.strip()]
        if len(annotations)==1 and int(annotations[0][0])==class_id:candidates.append((row,annotations))
    row,annotations=sorted(candidates,key=lambda pair:pair[0]['key'])[0]
    x=13.2+(i%2)*4.55;y=20.54+(i//2)*3.77
    shape(x,y,3.91,3.50,INK,None,MSO_SHAPE.ROUNDED_RECTANGLE)
    shape(x+.10,y+.10,3.71,3.30,WHITE,None,MSO_SHAPE.ROUNDED_RECTANGLE)
    shape(x+1.48,y+.13,.95,.055,INK,None,MSO_SHAPE.ROUNDED_RECTANGLE)
    px,py,pw,ph=x+.21,y+.36,3.49,2.26
    picture=s.shapes.add_picture(str(ROOT/row['image']),Inches(px),Inches(py),width=Inches(pw),height=Inches(ph))
    with Image.open(ROOT/row['image']) as photo:ratio=photo.width/photo.height
    target=pw/ph;left=top=0.;spanx=spany=1.
    if ratio>target:
        left=(1-target/ratio)/2;spanx=1-2*left;picture.crop_left=picture.crop_right=left
    else:
        top=(1-ratio/target)/2;spany=1-2*top;picture.crop_top=picture.crop_bottom=top
    for values in annotations:
        _,cx,cy,bw,bh=map(float,values)
        x1=max(0.,(cx-bw/2-left)/spanx);x2=min(1.,(cx+bw/2-left)/spanx)
        y1=max(0.,(cy-bh/2-top)/spany);y2=min(1.,(cy+bh/2-top)/spany)
        if x2>x1 and y2>y1:shape(px+x1*pw,py+y1*ph,(x2-x1)*pw,(y2-y1)*ph,None,'FFE57F',width=2.0)
    text(crop,x+.16,y+2.65,3.59,.39,20,GREEN,True,PP_ALIGN.CENTER)
    text(disease.replace('_',' ').title(),x+.16,y+3.09,3.59,.30,17,MUTED,False,PP_ALIGN.CENTER)
    photos.append(dict(crop=crop,class_name=names[str(class_id)],image=row['image'],label=row['label'],ground_truth=True))
text('Labelled validation examples • yellow boxes are ground truth',12.95,28.00,8.9,.25,16,MUTED,False,PP_ALIGN.CENTER)

# Reference footer: technologies on the left, summary on the right.
shape(.7,28.65,W-1.4,.025,CYAN)
shape(11.5,29.04,.02,3.10,LINE)
text('TECHNOLOGIES USED',1.05,28.91,10.0,.5,26,BLUE,True,PP_ALIGN.CENTER)
text('SUMMARY',12.0,28.91,10.0,.5,26,BLUE,True,PP_ALIGN.CENTER)
for i,(label,subtitle) in enumerate([('Python','Data & experiment tools'),('PyTorch','Detector implementations'),('Ultralytics','Historical YOLOv8n'),('Colab + Drive','GPU runs & checkpoints')]):
    x=1.15+(i%2)*5.08;y=29.72+(i//2)*1.12
    shape(x,y,4.77,.89,PALE,None,MSO_SHAPE.ROUNDED_RECTANGLE)
    text(label,x+.16,y+.08,4.45,.42,24,GREEN,True)
    text(subtitle,x+.16,y+.56,4.45,.28,18,MUTED)
text('The revised study contains four crops and 24 classes. The historical round-40 YOLO baseline is retained. '
     'The other three detectors will continue toward round 30. Final conclusions require completed runs, '
     'common evaluation and measured deployment cost.',12.15,29.73,9.67,2.2,24)
text('NEXT: complete training • evaluate retained YOLO on 24 classes • compare accuracy, speed and size',
     1.05,32.32,W-2.1,.37,21,BLUE,True,PP_ALIGN.CENTER)
text('DATA: 14,514 train / 4,157 validation / 2,091 test  •  3 simulated clients  •  Progress snapshot: 10 October 2026',
     1.05,32.84,W-2.1,.24,16,MUTED,False,PP_ALIGN.CENTER)

prs.core_properties.title='Agri Vision — revised FYP seminar poster'
prs.core_properties.author='Hannan Rehman, Fuzail Talpur, Ali'
prs.core_properties.subject='A1 portrait; stacked reference layout; current four-crop detection study'
prs.save(OUT/'Agri_Vision_Seminar_Poster_A1.pptx')
source=dict(format='A1 portrait, 594 × 841 mm',layout_reference='C:/Users/Hussaini Enterprise/Pictures/Screenshots/Screenshot 2026-10-09 115847.png',
            comparison_metrics='reports/comparison/main_run_review_20261010/validation_summary.csv',
            historical_metrics='reports/rounds_036_040_global_metrics.csv',
            dataset_examples=photos,examples_are_model_predictions=False,comparison_is_provisional=True,
            historical_yolo_common_24_class_evaluation='pending',fresh_yolo_training='on hold per user',
            university_logos='not supplied; no institutional branding invented',snapshot_date='2026-10-10')
(OUT/'sources.json').write_text(json.dumps(source,indent=2)+'\n',encoding='utf-8')
print(OUT/'Agri_Vision_Seminar_Poster_A1.pptx')
